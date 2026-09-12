import io

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.db.session import get_db
from app.models.archivo import Archivo
from app.models.archivo_diseno import ArchivoDiseno
from app.models.archivo_fila import ArchivoFila
from app.models.celda import Celda
from app.models.empresa import Empresa
from app.models.user import User
from app.routers.deps import get_current_user
from app.schemas.archivo import (
    ArchivoCreate,
    ArchivoFilaCreate,
    ArchivoFilaOut,
    ArchivoFilaUpdate,
    ArchivoOut,
    ArchivoUpdate,
    CeldaBulkUpdate,
    CeldaOut,
    CeldaUpdate,
    DisenoOut,
    DisenoUpdate,
)

router = APIRouter(prefix="/archivos", tags=["archivos"])


def _recalc_progreso(db: Session, archivo_id: str):
    total = db.query(func.count(Celda.id)).filter(Celda.archivo_id == archivo_id).scalar() or 0
    rev = (
        db.query(func.count(Celda.id))
        .filter(Celda.archivo_id == archivo_id, Celda.revisado.is_(True))
        .scalar()
        or 0
    )
    prog = int(rev * 100 / total) if total else 0
    arch = db.get(Archivo, archivo_id)
    if arch:
        arch.progreso = prog
        db.commit()
    return prog, total, rev


def _ensure_filas(db: Session, archivo: Archivo) -> list[ArchivoFila]:
    filas = (
        db.query(ArchivoFila)
        .filter(ArchivoFila.archivo_id == archivo.id)
        .order_by(ArchivoFila.orden)
        .all()
    )
    if filas:
        return filas
    # default: 5 empty rows, no data — months work per row, empresa picked later
    for idx in range(5):
        db.add(
            ArchivoFila(
                archivo_id=archivo.id,
                empresa_id=None,
                nombre_snapshot="",
                orden=idx,
            )
        )
    db.commit()
    return (
        db.query(ArchivoFila)
        .filter(ArchivoFila.archivo_id == archivo.id)
        .order_by(ArchivoFila.orden)
        .all()
    )


def _ensure_celdas(db: Session, archivo: Archivo, filas: list[ArchivoFila]):
    # create missing celdas per fila — months work with or without empresa
    for fila in filas:
        cnt = (
            db.query(func.count(Celda.id))
            .filter(Celda.archivo_id == archivo.id, Celda.fila_id == fila.id)
            .scalar()
            or 0
        )
        if cnt > 0:
            # ensure all years covered (for scale changes)
            for y in range(archivo.periodo_inicio, archivo.periodo_fin + 1):
                for m in range(1, 13):
                    if (
                        not db.query(Celda)
                        .filter(
                            Celda.archivo_id == archivo.id,
                            Celda.fila_id == fila.id,
                            Celda.anio == y,
                            Celda.mes == m,
                        )
                        .first()
                    ):
                        db.add(
                            Celda(
                                archivo_id=archivo.id,
                                fila_id=fila.id,
                                empresa_id=fila.empresa_id,
                                anio=y,
                                mes=m,
                            )
                        )
        else:
            for y in range(archivo.periodo_inicio, archivo.periodo_fin + 1):
                for m in range(1, 13):
                    db.add(
                        Celda(
                            archivo_id=archivo.id,
                            fila_id=fila.id,
                            empresa_id=fila.empresa_id,
                            anio=y,
                            mes=m,
                        )
                    )
    db.commit()


def _sync_escala(db: Session, archivo: Archivo, old_start: int, old_end: int):
    # delete out-of-range
    db.query(Celda).filter(
        Celda.archivo_id == archivo.id,
        (Celda.anio < archivo.periodo_inicio) | (Celda.anio > archivo.periodo_fin),
    ).delete(synchronize_session=False)
    # add missing years for each fila
    filas = db.query(ArchivoFila).filter(ArchivoFila.archivo_id == archivo.id).all()
    for fila in filas:
        for y in range(archivo.periodo_inicio, archivo.periodo_fin + 1):
            if y < old_start or y > old_end:
                for m in range(1, 13):
                    if (
                        not db.query(Celda)
                        .filter(
                            Celda.archivo_id == archivo.id,
                            Celda.fila_id == fila.id,
                            Celda.anio == y,
                            Celda.mes == m,
                        )
                        .first()
                    ):
                        db.add(
                            Celda(
                                archivo_id=archivo.id,
                                fila_id=fila.id,
                                empresa_id=fila.empresa_id,
                                anio=y,
                                mes=m,
                            )
                        )
    db.commit()


# ---------- Archivos ----------
@router.get("", response_model=list[ArchivoOut])
def list_archivos(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    return (
        db.query(Archivo)
        .filter(Archivo.user_id == user.id)
        .order_by(Archivo.updated_at.desc())
        .all()
    )


@router.post("", response_model=ArchivoOut)
def create(
    data: ArchivoCreate, db: Session = Depends(get_db), user: User = Depends(get_current_user)
):
    a = Archivo(
        user_id=user.id,
        titulo=data.titulo,
        tipo_revision=data.tipo_revision,
        periodo_inicio=data.periodo_inicio,
        periodo_fin=data.periodo_fin,
        personal_count=data.personal_count or 2,
    )
    db.add(a)
    db.commit()
    db.refresh(a)
    filas = _ensure_filas(db, a)
    if filas:
        _ensure_celdas(db, a, filas)
        _recalc_progreso(db, a.id)
        db.refresh(a)
    return a


@router.get("/{aid}", response_model=ArchivoOut)
def get_one(aid: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    a = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not a:
        raise HTTPException(404, "No encontrado")
    return a


@router.put("/{aid}", response_model=ArchivoOut)
def update(
    aid: str,
    data: ArchivoUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    a = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not a:
        raise HTTPException(404, "No encontrado")
    old_start, old_end = a.periodo_inicio, a.periodo_fin
    if data.titulo is not None:
        a.titulo = data.titulo
    if data.tipo_revision is not None:
        a.tipo_revision = data.tipo_revision
    if data.personal_count is not None:
        if data.personal_count < 1 or data.personal_count > 10:
            raise HTTPException(400, "personal_count debe ser 1-10")
        a.personal_count = data.personal_count
    escala_changed = False
    if data.periodo_inicio is not None and data.periodo_inicio != a.periodo_inicio:
        a.periodo_inicio = data.periodo_inicio
        escala_changed = True
    if data.periodo_fin is not None and data.periodo_fin != a.periodo_fin:
        a.periodo_fin = data.periodo_fin
        escala_changed = True
    if a.periodo_inicio > a.periodo_fin:
        raise HTTPException(400, "Rango de años inválido")
    if escala_changed and (a.periodo_fin - a.periodo_inicio) > 20:
        raise HTTPException(400, "Rango demasiado amplio (máx 20 años)")
    db.commit()
    db.refresh(a)
    if escala_changed:
        _sync_escala(db, a, old_start, old_end)
    return a


SECCIONES_DISENO = ("hoja", "tabla")


def _check_archivo(aid: str, db: Session, user: User) -> Archivo:
    a = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not a:
        raise HTTPException(404, "No encontrado")
    return a


@router.get("/{aid}/diseno", response_model=list[DisenoOut])
def get_diseno(aid: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    _check_archivo(aid, db, user)
    return (
        db.query(ArchivoDiseno)
        .filter(ArchivoDiseno.archivo_id == aid, ArchivoDiseno.seccion.in_(SECCIONES_DISENO))
        .all()
    )


@router.put("/{aid}/diseno", response_model=DisenoOut)
def put_diseno(
    aid: str,
    data: DisenoUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    _check_archivo(aid, db, user)
    if data.seccion not in SECCIONES_DISENO:
        raise HTTPException(400, "seccion debe ser 'hoja' o 'tabla'")
    if not isinstance(data.payload, dict) or len(data.payload) > 500:
        raise HTTPException(400, "payload inválido")
    row = (
        db.query(ArchivoDiseno)
        .filter(ArchivoDiseno.archivo_id == aid, ArchivoDiseno.seccion == data.seccion)
        .first()
    )
    if not row:
        row = ArchivoDiseno(archivo_id=aid, seccion=data.seccion, payload=data.payload)
        db.add(row)
    else:
        row.payload = data.payload
    db.commit()
    db.refresh(row)
    return row


@router.delete("/{aid}")
def delete(aid: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    a = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not a:
        raise HTTPException(404, "No encontrado")
    db.delete(a)
    db.commit()
    return {"ok": True}


@router.post("/{aid}/duplicate", response_model=ArchivoOut)
def duplicate(aid: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    orig = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not orig:
        raise HTTPException(404, "No encontrado")
    a = Archivo(
        user_id=user.id,
        titulo=orig.titulo + " (copia)",
        tipo_revision=orig.tipo_revision,
        periodo_inicio=orig.periodo_inicio,
        periodo_fin=orig.periodo_fin,
    )
    db.add(a)
    db.commit()
    db.refresh(a)
    filas = (
        db.query(ArchivoFila)
        .filter(ArchivoFila.archivo_id == orig.id)
        .order_by(ArchivoFila.orden)
        .all()
    )
    for f in filas:
        nf = ArchivoFila(
            archivo_id=a.id,
            empresa_id=f.empresa_id,
            nombre_snapshot=f.nombre_snapshot,
            orden=f.orden,
            responsable=f.responsable,
            observacion=f.observacion,
        )
        db.add(nf)
    db.commit()
    # copy each row with its own month cells (fresh revision: unchecked)
    new_filas = (
        db.query(ArchivoFila)
        .filter(ArchivoFila.archivo_id == a.id)
        .order_by(ArchivoFila.orden)
        .all()
    )
    old_celdas = db.query(Celda).filter(Celda.archivo_id == orig.id).all()
    old_by_fila: dict = {}
    for c in old_celdas:
        old_by_fila.setdefault(c.fila_id, []).append(c)
    old_by_orden = {f.orden: f.id for f in filas}
    for nf in new_filas:
        for c in old_by_fila.get(old_by_orden.get(nf.orden), []):
            db.add(
                Celda(
                    archivo_id=a.id,
                    fila_id=nf.id,
                    empresa_id=nf.empresa_id,
                    anio=c.anio,
                    mes=c.mes,
                    revisado=False,
                    responsable=c.responsable,
                    color=c.color,
                    style=c.style,
                )
            )
    db.commit()
    _recalc_progreso(db, a.id)
    # copiar el diseño visual (hoja + tabla; los altos por fila se reasignan por orden,
    # las claves que no son filas —ej. "header"— se preservan verbatim)
    for d in db.query(ArchivoDiseno).filter(ArchivoDiseno.archivo_id == orig.id).all():
        payload = d.payload if isinstance(d.payload, dict) else {}
        if d.seccion == "tabla" and isinstance(payload.get("rows"), dict):
            new_by_orden = {f.orden: f.id for f in new_filas}
            old_ids = {f.id for f in filas}
            rows = {k: v for k, v in payload["rows"].items() if k not in old_ids}
            for f in filas:
                v = payload["rows"].get(f.id)
                if v is not None and f.orden in new_by_orden:
                    rows[new_by_orden[f.orden]] = v
            payload = {**payload, "rows": rows}
        db.add(ArchivoDiseno(archivo_id=a.id, seccion=d.seccion, payload=payload))
    db.commit()
    db.refresh(a)
    return a


# ---------- Filas ----------
@router.get("/{aid}/filas", response_model=list[ArchivoFilaOut])
def list_filas(aid: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    arch = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not arch:
        raise HTTPException(404, "No encontrado")
    filas = _ensure_filas(db, arch)
    _ensure_celdas(db, arch, filas)
    return (
        db.query(ArchivoFila)
        .filter(ArchivoFila.archivo_id == aid)
        .order_by(ArchivoFila.orden)
        .all()
    )


@router.post("/{aid}/filas", response_model=ArchivoFilaOut)
def create_fila(
    aid: str,
    data: ArchivoFilaCreate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    arch = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not arch:
        raise HTTPException(404, "No encontrado")
    nombre = None
    empresa_id = data.empresa_id
    if empresa_id:
        emp = db.query(Empresa).filter(Empresa.id == empresa_id, Empresa.user_id == user.id).first()
        if not emp:
            raise HTTPException(404, "Empresa no encontrada")
        nombre = emp.nombre
    elif data.nombre:
        nombre = data.nombre.strip()
    else:
        raise HTTPException(400, "Debe enviar empresa_id o nombre")
    max_orden = db.query(func.max(ArchivoFila.orden)).filter(ArchivoFila.archivo_id == aid).scalar()
    orden = (max_orden + 1) if max_orden is not None else 0
    fila = ArchivoFila(archivo_id=aid, empresa_id=empresa_id, nombre_snapshot=nombre, orden=orden)
    db.add(fila)
    db.commit()
    db.refresh(fila)
    # every row gets its own month cells, with or without empresa
    for y in range(arch.periodo_inicio, arch.periodo_fin + 1):
        for m in range(1, 13):
            db.add(Celda(archivo_id=aid, fila_id=fila.id, empresa_id=empresa_id, anio=y, mes=m))
    db.commit()
    return fila


@router.put("/{aid}/filas/{fid}", response_model=ArchivoFilaOut)
def update_fila(
    aid: str,
    fid: str,
    data: ArchivoFilaUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    arch = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not arch:
        raise HTTPException(404, "No encontrado")
    fila = (
        db.query(ArchivoFila).filter(ArchivoFila.id == fid, ArchivoFila.archivo_id == aid).first()
    )
    if not fila:
        raise HTTPException(404, "Fila no encontrada")
    # empresa is an optional directory link: it only sets the label, never touches month cells
    if data.empresa_id is not None:
        if data.empresa_id == "":
            fila.empresa_id = None
            if data.nombre:
                fila.nombre_snapshot = data.nombre.strip()
        else:
            emp = (
                db.query(Empresa)
                .filter(Empresa.id == data.empresa_id, Empresa.user_id == user.id)
                .first()
            )
            if not emp:
                raise HTTPException(404, "Empresa no encontrada")
            fila.empresa_id = emp.id
            fila.nombre_snapshot = emp.nombre
    elif data.nombre is not None:
        # display label only — empresas must be pre-registered, never auto-created
        nombre = data.nombre.strip()
        if not nombre:
            raise HTTPException(400, "Nombre vacío")
        fila.nombre_snapshot = nombre
    if data.orden is not None:
        fila.orden = data.orden
    if data.responsable is not None:
        fila.responsable = data.responsable.strip() or None
    if data.observacion is not None:
        fila.observacion = data.observacion.strip() or None
    db.commit()
    db.refresh(fila)
    return fila


@router.delete("/{aid}/filas/{fid}")
def delete_fila(
    aid: str, fid: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)
):
    arch = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not arch:
        raise HTTPException(404, "No encontrado")
    fila = (
        db.query(ArchivoFila).filter(ArchivoFila.id == fid, ArchivoFila.archivo_id == aid).first()
    )
    if not fila:
        raise HTTPException(404, "Fila no encontrada")
    # month cells belong to the row: deleting it always removes its own cells
    db.query(Celda).filter(Celda.archivo_id == aid, Celda.fila_id == fid).delete(
        synchronize_session=False
    )
    db.delete(fila)
    db.commit()
    # reindex orden
    filas = (
        db.query(ArchivoFila)
        .filter(ArchivoFila.archivo_id == aid)
        .order_by(ArchivoFila.orden)
        .all()
    )
    for idx, f in enumerate(filas):
        f.orden = idx
    db.commit()
    # podar el alto guardado de la fila eliminada
    dis = (
        db.query(ArchivoDiseno)
        .filter(ArchivoDiseno.archivo_id == aid, ArchivoDiseno.seccion == "tabla")
        .first()
    )
    if dis and isinstance(dis.payload, dict) and fid in (dis.payload.get("rows") or {}):
        payload = dict(dis.payload)
        rows = dict(payload.get("rows") or {})
        rows.pop(fid, None)
        payload["rows"] = rows
        dis.payload = payload
        db.commit()
    return {"ok": True}


@router.post("/{aid}/filas/reorder")
def reorder_filas(
    aid: str, payload: dict, db: Session = Depends(get_db), user: User = Depends(get_current_user)
):
    arch = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not arch:
        raise HTTPException(404, "No encontrado")
    order: list[str] = payload.get("order", [])
    for idx, fid in enumerate(order):
        db.query(ArchivoFila).filter(ArchivoFila.id == fid, ArchivoFila.archivo_id == aid).update(
            {"orden": idx}
        )
    db.commit()
    return {"ok": True}


# ---------- Celdas ----------
@router.get("/{aid}/celdas", response_model=list[CeldaOut])
def list_celdas(aid: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    arch = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not arch:
        raise HTTPException(404, "No encontrado")
    filas = _ensure_filas(db, arch)
    _ensure_celdas(db, arch, filas)
    return db.query(Celda).filter(Celda.archivo_id == aid).all()


@router.put("/../celdas/{cid}", response_model=CeldaOut)
def update_celda(
    cid: str,
    data: CeldaUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    c = db.get(Celda, cid)
    if not c:
        raise HTTPException(404, "No encontrado")
    arch = db.get(Archivo, c.archivo_id)
    if not arch or arch.user_id != user.id:
        raise HTTPException(403, "No autorizado")
    if data.revisado is not None:
        c.revisado = data.revisado
    if data.responsable is not None:
        c.responsable = data.responsable
    if data.observacion is not None:
        c.observacion = data.observacion
    if data.color is not None:
        c.color = data.color or None
    if data.style is not None:
        c.style = data.style
    db.commit()
    db.refresh(c)
    _recalc_progreso(db, arch.id)
    return c


# alias for /celdas/{id}
@router.put("/celdas/{cid}", response_model=CeldaOut, include_in_schema=False)
def update_celda_alias(
    cid: str,
    data: CeldaUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    return update_celda(cid, data, db, user)


@router.post("/{aid}/celdas/bulk", response_model=list[CeldaOut])
def bulk_update(
    aid: str,
    data: CeldaBulkUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    arch = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not arch:
        raise HTTPException(404, "No encontrado")
    out = []
    for cid in data.ids:
        c = db.get(Celda, cid)
        if not c or c.archivo_id != aid:
            continue
        if data.color is not None:
            c.color = data.color
        if data.style is not None:
            # merge style
            cur = c.style or {}
            cur.update(data.style)
            c.style = cur
        if data.revisado is not None:
            c.revisado = data.revisado
        out.append(c)
    db.commit()
    for c in out:
        db.refresh(c)
    _recalc_progreso(db, aid)
    return out


@router.get("/{aid}/stats")
def stats(aid: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    arch = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not arch:
        raise HTTPException(404, "No encontrado")
    total = db.query(func.count(Celda.id)).filter(Celda.archivo_id == aid).scalar() or 0
    rev = (
        db.query(func.count(Celda.id))
        .filter(Celda.archivo_id == aid, Celda.revisado.is_(True))
        .scalar()
        or 0
    )
    filas = _ensure_filas(db, arch)
    _ensure_celdas(db, arch, filas)
    # count fully-reviewed rows (every row counts, with or without empresa)
    revisadas = 0
    for fila in filas:
        t = (
            db.query(func.count(Celda.id))
            .filter(Celda.archivo_id == aid, Celda.fila_id == fila.id)
            .scalar()
            or 0
        )
        r = (
            db.query(func.count(Celda.id))
            .filter(
                Celda.archivo_id == aid,
                Celda.fila_id == fila.id,
                Celda.revisado.is_(True),
            )
            .scalar()
            or 0
        )
        if t > 0 and t == r:
            revisadas += 1
    total_filas = len(filas)
    pendientes = total_filas - revisadas if total_filas else 0
    prog = int(rev * 100 / total) if total else 0
    return {
        "total": total_filas,
        "revisadas": revisadas,
        "pendientes": pendientes,
        "progreso": prog,
        "celdas_total": total,
        "celdas_revisadas": rev,
    }


def _missing_fields(filas: list[ArchivoFila]) -> list[dict]:
    """Required fields for save/export: every row needs a registered empresa."""
    missing = []
    for idx, fila in enumerate(sorted(filas, key=lambda f: f.orden), 1):
        if not fila.empresa_id:
            missing.append(
                {"fila": idx, "id": fila.id, "campo": "empresa", "nombre": fila.nombre_snapshot}
            )
    return missing


@router.get("/{aid}/export")
def export_file(
    aid: str,
    format: str = "pdf",
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    arch = db.query(Archivo).filter(Archivo.id == aid, Archivo.user_id == user.id).first()
    if not arch:
        raise HTTPException(404, "No encontrado")
    filas = _ensure_filas(db, arch)
    _ensure_celdas(db, arch, filas)
    filas = sorted(filas, key=lambda f: f.orden)
    missing = _missing_fields(filas)
    if missing:
        raise HTTPException(
            status_code=422,
            detail={"message": "Faltan campos requeridos para exportar", "faltantes": missing},
        )
    celdas = db.query(Celda).filter(Celda.archivo_id == aid).all()
    cmap = {(c.fila_id, c.anio, c.mes): c for c in celdas}

    if format == "excel":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        wb = Workbook()
        ws = wb.active
        ws.title = "Revision"
        header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=9)
        thin = Side(style="thin", color="E2E8F0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        years = list(range(arch.periodo_inicio, arch.periodo_fin + 1))
        meses = ["E", "F", "M", "A", "M", "J", "J", "A", "S", "O", "N", "D"]
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
        ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
        ws.cell(row=1, column=1, value="N.°").font = header_font
        ws.cell(row=1, column=1).fill = header_fill
        ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=1, column=2, value="Empresa").font = header_font
        ws.cell(row=1, column=2).fill = header_fill
        ws.cell(row=1, column=2).alignment = Alignment(horizontal="center", vertical="center")
        col = 3
        for y in years:
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 11)
            c = ws.cell(row=1, column=col, value=str(y))
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center")
            col += 12
        col = 3
        for _ in years:
            for m in meses:
                c = ws.cell(row=2, column=col, value=m)
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(horizontal="center")
                col += 1
        for hdr in ("Responsable", "Observaciones"):
            ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
            c = ws.cell(row=1, column=col, value=hdr)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            col += 1
        r = 3
        for idx, fila in enumerate(filas, 1):
            ws.cell(row=r, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=2, value=fila.nombre_snapshot)
            col = 3
            for y in years:
                for m_idx in range(1, 13):
                    c = cmap.get((fila.id, y, m_idx))
                    v = "☑" if c and c.revisado else "☐"
                    cell = ws.cell(row=r, column=col, value=v)
                    cell.alignment = Alignment(horizontal="center")
                    if c and c.revisado:
                        colr = (c.color or "6366F1").lstrip("#")
                        cell.font = Font(color=colr)
                        # style bold
                        if c.style and c.style.get("bold"):
                            cell.font = Font(color=colr, bold=True)
                    col += 1
            ws.cell(row=r, column=col, value=fila.responsable or "")
            ws.cell(row=r, column=col + 1, value=fila.observacion or "")
            col += 2
            r += 1
        for row in ws.iter_rows(min_row=1, max_row=r - 1, max_col=col - 1):
            for cell in row:
                cell.border = border
        ws.freeze_panes = "C3"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return StreamingResponse(
            bio,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=revision-{aid}.xlsx"},
        )
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        bio = io.BytesIO()
        doc = SimpleDocTemplate(
            bio,
            pagesize=landscape(A4),
            leftMargin=10 * mm,
            rightMargin=10 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm,
        )
        styles = getSampleStyleSheet()
        story = []
        story.append(
            Paragraph(
                f"{arch.titulo or 'Revisión sin título'} — {arch.periodo_inicio}-{arch.periodo_fin} | {len(filas)} Empresas",
                styles["Title"],
            )
        )
        story.append(Spacer(1, 6))
        years = list(range(arch.periodo_inicio, arch.periodo_fin + 1))
        meses = ["E", "F", "M", "A", "M", "J", "J", "A", "S", "O", "N", "D"]
        flat_header = ["N.°", "Empresa"]
        for y in years:
            for m in meses:
                flat_header.append(f"{y}-{m}")
        flat_header += ["Responsable", "Observaciones"]
        table_data = [flat_header]
        for idx, fila in enumerate(filas, 1):
            row = [str(idx), fila.nombre_snapshot]
            for y in years:
                for m_idx in range(1, 13):
                    c = cmap.get((fila.id, y, m_idx))
                    row.append("☑" if c and c.revisado else "☐")
            row += [fila.responsable or "", fila.observacion or ""]
            table_data.append(row)
        n_meses = len(flat_header) - 4
        col_widths = [12 * mm, 30 * mm] + [8 * mm] * n_meses + [20 * mm, 28 * mm]
        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        style = TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6366f1")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, 0), 6),
                ("FONTSIZE", (0, 1), (-1, -1), 5),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ]
        )
        t.setStyle(style)
        story.append(t)
        doc.build(story)
        bio.seek(0)
        return StreamingResponse(
            bio,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename=revision-{aid}.pdf"},
        )
