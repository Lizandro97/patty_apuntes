import io
from datetime import UTC, datetime

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.db.session import get_db
from app.export.filenames import export_filename
from app.models.cell import Cell
from app.models.company import Company
from app.models.record import Record
from app.models.record_layout import RecordLayout
from app.models.record_row import RecordRow
from app.models.user import User
from app.routers.deps import get_current_user
from app.schemas.record import (
    CellBulkUpdate,
    CellOut,
    CellUpdate,
    LayoutOut,
    LayoutUpdate,
    RecordCreate,
    RecordOut,
    RecordRowCreate,
    RecordRowOut,
    RecordRowUpdate,
    RecordUpdate,
)
from app.sync.engine import touch_record
from app.validation.rules import (
    find_rows_missing_company,
    sanitize_staff_names,
    validate_scale,
    validate_staff,
    validate_staff_names,
)

router = APIRouter(prefix="/records", tags=["records"])


def _recalc_progress(db: Session, record_id: str):
    total = db.query(func.count(Cell.id)).filter(Cell.record_id == record_id).scalar() or 0
    rev = (
        db.query(func.count(Cell.id))
        .filter(Cell.record_id == record_id, Cell.reviewed.is_(True))
        .scalar()
        or 0
    )
    prog = int(rev * 100 / total) if total else 0
    record = db.get(Record, record_id)
    if record:
        record.progress = prog
        db.commit()
    return prog, total, rev


def _ensure_rows(db: Session, record: Record) -> list[RecordRow]:
    rows = (
        db.query(RecordRow)
        .filter(RecordRow.record_id == record.id)
        .order_by(RecordRow.position)
        .all()
    )
    if rows:
        return rows
    # default: 5 empty rows, no data — months work per row, empresa picked later
    for idx in range(5):
        db.add(
            RecordRow(
                record_id=record.id,
                company_id=None,
                name_snapshot="",
                position=idx,
            )
        )
    db.commit()
    return (
        db.query(RecordRow)
        .filter(RecordRow.record_id == record.id)
        .order_by(RecordRow.position)
        .all()
    )


def _ensure_cells(db: Session, record: Record, rows: list[RecordRow]):
    # create missing cells per row — months work with or without empresa
    for row in rows:
        cnt = (
            db.query(func.count(Cell.id))
            .filter(Cell.record_id == record.id, Cell.row_id == row.id)
            .scalar()
            or 0
        )
        if cnt > 0:
            # ensure all years covered (for scale changes)
            for y in range(record.period_start, record.period_end + 1):
                for m in range(1, 13):
                    if (
                        not db.query(Cell)
                        .filter(
                            Cell.record_id == record.id,
                            Cell.row_id == row.id,
                            Cell.year == y,
                            Cell.month == m,
                        )
                        .first()
                    ):
                        db.add(
                            Cell(
                                record_id=record.id,
                                row_id=row.id,
                                company_id=row.company_id,
                                year=y,
                                month=m,
                            )
                        )
        else:
            for y in range(record.period_start, record.period_end + 1):
                for m in range(1, 13):
                    db.add(
                        Cell(
                            record_id=record.id,
                            row_id=row.id,
                            company_id=row.company_id,
                            year=y,
                            month=m,
                        )
                    )
    db.commit()


def _sync_scale(db: Session, record: Record, old_start: int, old_end: int):
    # delete out-of-range
    db.query(Cell).filter(
        Cell.record_id == record.id,
        (Cell.year < record.period_start) | (Cell.year > record.period_end),
    ).delete(synchronize_session=False)
    # add missing years for each row
    rows = db.query(RecordRow).filter(RecordRow.record_id == record.id).all()
    for row in rows:
        for y in range(record.period_start, record.period_end + 1):
            if y < old_start or y > old_end:
                for m in range(1, 13):
                    if (
                        not db.query(Cell)
                        .filter(
                            Cell.record_id == record.id,
                            Cell.row_id == row.id,
                            Cell.year == y,
                            Cell.month == m,
                        )
                        .first()
                    ):
                        db.add(
                            Cell(
                                record_id=record.id,
                                row_id=row.id,
                                company_id=row.company_id,
                                year=y,
                                month=m,
                            )
                        )
    db.commit()


def _ensure_live(db: Session, record: Record) -> list[RecordRow]:
    """Ensure por defecto + touch solo si realmente creo filas/celdas.

    Los GET que materializan defaults (list_rows/cells/stats/export) tambien
    generan datos sync-visibles: sin touch el pull los perderia.
    """
    rows_before = (
        db.query(func.count(RecordRow.id)).filter(RecordRow.record_id == record.id).scalar() or 0
    )
    cells_before = db.query(func.count(Cell.id)).filter(Cell.record_id == record.id).scalar() or 0
    rows = _ensure_rows(db, record)
    _ensure_cells(db, record, rows)
    rows_after = (
        db.query(func.count(RecordRow.id)).filter(RecordRow.record_id == record.id).scalar() or 0
    )
    cells_after = db.query(func.count(Cell.id)).filter(Cell.record_id == record.id).scalar() or 0
    if rows_after != rows_before or cells_after != cells_before:
        touch_record(db, record)
    return (
        db.query(RecordRow)
        .filter(RecordRow.record_id == record.id)
        .order_by(RecordRow.position)
        .all()
    )


# ---------- Archivos ----------
@router.get("", response_model=list[RecordOut])
def list_records(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    return (
        db.query(Record)
        .filter(Record.user_id == user.id, Record.deleted_at.is_(None))
        .order_by(Record.updated_at.desc())
        .all()
    )


@router.post("", response_model=RecordOut)
def create(
    data: RecordCreate, db: Session = Depends(get_db), user: User = Depends(get_current_user)
):
    a = Record(
        user_id=user.id,
        title=data.title,
        review_type=data.review_type,
        period_start=data.period_start,
        period_end=data.period_end,
        staff_count=data.staff_count or 2,
        staff_names=sanitize_staff_names(data.staff_names),
    )
    db.add(a)
    db.commit()
    db.refresh(a)
    rows = _ensure_rows(db, a)
    if rows:
        _ensure_cells(db, a, rows)
        _recalc_progress(db, a.id)
        db.refresh(a)
    touch_record(db, a)
    return a


@router.get("/{record_id}", response_model=RecordOut)
def get_one(record_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    a = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not a:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    return a


@router.put("/{record_id}", response_model=RecordOut)
def update(
    record_id: str,
    data: RecordUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    a = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not a:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    old_start, old_end = a.period_start, a.period_end
    if data.title is not None:
        a.title = data.title
    if data.review_type is not None:
        a.review_type = data.review_type
    if data.staff_count is not None:
        err = validate_staff(data.staff_count)
        if err:
            raise HTTPException(400, {"code": err["code"], "message": err["message"]})
        a.staff_count = data.staff_count
    if data.staff_names is not None:
        err = validate_staff_names(data.staff_names)
        if err:
            raise HTTPException(400, {"code": err["code"], "message": err["message"]})
        a.staff_names = sanitize_staff_names(data.staff_names)
    scale_changed = False
    if data.period_start is not None and data.period_start != a.period_start:
        a.period_start = data.period_start
        scale_changed = True
    if data.period_end is not None and data.period_end != a.period_end:
        a.period_end = data.period_end
        scale_changed = True
    err = validate_scale(a.period_start, a.period_end)
    if err and (scale_changed or err["code"] == "INVALID_SCALE"):
        raise HTTPException(400, {"code": err["code"], "message": err["message"]})
    db.commit()
    db.refresh(a)
    if scale_changed:
        _sync_scale(db, a, old_start, old_end)
    touch_record(db, a)
    return a


LAYOUT_SECTIONS = ("sheet", "table")


def _require_record(record_id: str, db: Session, user: User) -> Record:
    a = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not a:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    return a


@router.get("/{record_id}/layout", response_model=list[LayoutOut])
def get_layout(
    record_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)
):
    _require_record(record_id, db, user)
    return (
        db.query(RecordLayout)
        .filter(RecordLayout.record_id == record_id, RecordLayout.section.in_(LAYOUT_SECTIONS))
        .all()
    )


@router.put("/{record_id}/layout", response_model=LayoutOut)
def put_layout(
    record_id: str,
    data: LayoutUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    record = _require_record(record_id, db, user)
    if data.section not in LAYOUT_SECTIONS:
        raise HTTPException(
            400, {"code": "SECTION_INVALID", "message": "section debe ser 'sheet' o 'table'"}
        )
    if not isinstance(data.payload, dict) or len(data.payload) > 500:
        raise HTTPException(400, {"code": "PAYLOAD_INVALID", "message": "Invalid payload"})
    row = (
        db.query(RecordLayout)
        .filter(RecordLayout.record_id == record_id, RecordLayout.section == data.section)
        .first()
    )
    if not row:
        row = RecordLayout(record_id=record_id, section=data.section, payload=data.payload)
        db.add(row)
    else:
        row.payload = data.payload
    db.commit()
    db.refresh(row)
    touch_record(db, record)
    return row


@router.delete("/{record_id}")
def delete(record_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    a = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not a:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    # Borrado logico (Fase 2): viaja como tombstone en pull, no se pierde en sync.
    a.deleted_at = datetime.now(UTC)
    touch_record(db, a)
    return {"ok": True}


@router.post("/{record_id}/duplicate", response_model=RecordOut)
def duplicate(
    record_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)
):
    orig = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not orig:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    a = Record(
        user_id=user.id,
        title=orig.title + " (copia)",
        review_type=orig.review_type,
        period_start=orig.period_start,
        period_end=orig.period_end,
    )
    db.add(a)
    db.commit()
    db.refresh(a)
    rows = (
        db.query(RecordRow)
        .filter(RecordRow.record_id == orig.id)
        .order_by(RecordRow.position)
        .all()
    )
    for f in rows:
        new_row = RecordRow(
            record_id=a.id,
            company_id=f.company_id,
            name_snapshot=f.name_snapshot,
            position=f.position,
            assignee=f.assignee,
            note=f.note,
        )
        db.add(new_row)
    db.commit()
    # copy each row with its own month cells (fresh revision: unchecked)
    new_rows = (
        db.query(RecordRow).filter(RecordRow.record_id == a.id).order_by(RecordRow.position).all()
    )
    old_cells = db.query(Cell).filter(Cell.record_id == orig.id).all()
    old_by_row: dict = {}
    for c in old_cells:
        old_by_row.setdefault(c.row_id, []).append(c)
    old_by_position = {f.position: f.id for f in rows}
    for new_row in new_rows:
        for c in old_by_row.get(old_by_position.get(new_row.position), []):
            db.add(
                Cell(
                    record_id=a.id,
                    row_id=new_row.id,
                    company_id=new_row.company_id,
                    year=c.year,
                    month=c.month,
                    reviewed=False,
                    assignee=c.assignee,
                    color=c.color,
                    style=c.style,
                )
            )
    db.commit()
    _recalc_progress(db, a.id)
    # Copy the visual design (sheet + table; per-row heights are reassigned by
    # position, non-row keys —e.g. "header"— are preserved verbatim)
    for d in db.query(RecordLayout).filter(RecordLayout.record_id == orig.id).all():
        payload = d.payload if isinstance(d.payload, dict) else {}
        if d.section == "table" and isinstance(payload.get("rows"), dict):
            new_by_position = {f.position: f.id for f in new_rows}
            old_ids = {f.id for f in rows}
            remapped = {k: v for k, v in payload["rows"].items() if k not in old_ids}
            for f in rows:
                v = payload["rows"].get(f.id)
                if v is not None and f.position in new_by_position:
                    remapped[new_by_position[f.position]] = v
            payload = {**payload, "rows": remapped}
        db.add(RecordLayout(record_id=a.id, section=d.section, payload=payload))
    db.commit()
    db.refresh(a)
    touch_record(db, a)
    return a


# ---------- Filas ----------
@router.get("/{record_id}/rows", response_model=list[RecordRowOut])
def list_rows(
    record_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)
):
    record = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not record:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    return _ensure_live(db, record)


@router.post("/{record_id}/rows", response_model=RecordRowOut)
def create_row(
    record_id: str,
    data: RecordRowCreate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    record = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not record:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    name = None
    company_id = data.company_id
    if company_id:
        company = (
            db.query(Company).filter(Company.id == company_id, Company.user_id == user.id).first()
        )
        if not company:
            raise HTTPException(404, {"code": "COMPANY_NOT_FOUND", "message": "Company not found"})
        name = company.name
    elif data.name and data.name.strip():
        name = data.name.strip()
    else:
        # empty rows allowed (same as _ensure_rows defaults): they render the
        # "pick company" hint and are blocked by validation on save/export
        name = ""
        company_id = None
    max_position = (
        db.query(func.max(RecordRow.position)).filter(RecordRow.record_id == record_id).scalar()
    )
    position = (max_position + 1) if max_position is not None else 0
    row = RecordRow(
        record_id=record_id, company_id=company_id, name_snapshot=name, position=position
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    # every row gets its own month cells, with or without empresa
    for y in range(record.period_start, record.period_end + 1):
        for m in range(1, 13):
            db.add(Cell(record_id=record_id, row_id=row.id, company_id=company_id, year=y, month=m))
    db.commit()
    touch_record(db, record)
    return row


@router.put("/{record_id}/rows/{row_id}", response_model=RecordRowOut)
def update_row(
    record_id: str,
    row_id: str,
    data: RecordRowUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    record = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not record:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    row = (
        db.query(RecordRow).filter(RecordRow.id == row_id, RecordRow.record_id == record_id).first()
    )
    if not row:
        raise HTTPException(404, {"code": "ROW_NOT_FOUND", "message": "Row not found"})
    # empresa is an optional directory link: it only sets the label, never touches month cells
    if data.company_id is not None:
        if data.company_id == "":
            row.company_id = None
            if data.name:
                row.name_snapshot = data.name.strip()
        else:
            company = (
                db.query(Company)
                .filter(Company.id == data.company_id, Company.user_id == user.id)
                .first()
            )
            if not company:
                raise HTTPException(
                    404, {"code": "COMPANY_NOT_FOUND", "message": "Company not found"}
                )
            row.company_id = company.id
            row.name_snapshot = company.name
    elif data.name is not None:
        # display label only — companies must be pre-registered, never auto-created
        name = data.name.strip()
        if not name:
            raise HTTPException(400, {"code": "EMPTY_NAME", "message": "Empty name"})
        row.name_snapshot = name
    if data.position is not None:
        row.position = data.position
    if data.assignee is not None:
        row.assignee = data.assignee.strip() or None
    if data.note is not None:
        row.note = data.note.strip() or None
    db.commit()
    db.refresh(row)
    touch_record(db, record)
    return row


@router.delete("/{record_id}/rows/{row_id}")
def delete_row(
    record_id: str,
    row_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    record = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not record:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    row = (
        db.query(RecordRow).filter(RecordRow.id == row_id, RecordRow.record_id == record_id).first()
    )
    if not row:
        raise HTTPException(404, {"code": "ROW_NOT_FOUND", "message": "Row not found"})
    # month cells belong to the row: deleting it always removes its own cells
    db.query(Cell).filter(Cell.record_id == record_id, Cell.row_id == row_id).delete(
        synchronize_session=False
    )
    db.delete(row)
    db.commit()
    # reindex position
    rows = (
        db.query(RecordRow)
        .filter(RecordRow.record_id == record_id)
        .order_by(RecordRow.position)
        .all()
    )
    for idx, f in enumerate(rows):
        f.position = idx
    db.commit()
    # podar el alto guardado de la row eliminada
    dis = (
        db.query(RecordLayout)
        .filter(RecordLayout.record_id == record_id, RecordLayout.section == "table")
        .first()
    )
    if dis and isinstance(dis.payload, dict) and row_id in (dis.payload.get("rows") or {}):
        payload = dict(dis.payload)
        rows = dict(payload.get("rows") or {})
        rows.pop(row_id, None)
        payload["rows"] = rows
        dis.payload = payload
        db.commit()
    touch_record(db, record)
    return {"ok": True}


@router.post("/{record_id}/rows/reorder")
def reorder_rows(
    record_id: str,
    payload: dict,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    record = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not record:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    order: list[str] = payload.get("ids", payload.get("order", []))
    for idx, row_id in enumerate(order):
        db.query(RecordRow).filter(RecordRow.id == row_id, RecordRow.record_id == record_id).update(
            {"position": idx}
        )
    db.commit()
    touch_record(db, record)
    return {"ok": True}


# ---------- Celdas ----------
@router.get("/{record_id}/cells", response_model=list[CellOut])
def list_cells(
    record_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)
):
    record = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not record:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    _ensure_live(db, record)
    return db.query(Cell).filter(Cell.record_id == record_id).all()


@router.put("/../cells/{cell_id}", response_model=CellOut)
def update_cell(
    cell_id: str,
    data: CellUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    c = db.get(Cell, cell_id)
    if not c:
        raise HTTPException(404, {"code": "CELL_NOT_FOUND", "message": "Not found"})
    record = db.get(Record, c.record_id)
    if not record or record.user_id != user.id:
        raise HTTPException(403, {"code": "UNAUTHORIZED", "message": "Unauthorized"})
    if record.deleted_at is not None:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    if data.reviewed is not None:
        c.reviewed = data.reviewed
    if data.assignee is not None:
        c.assignee = data.assignee
    if data.note is not None:
        c.note = data.note
    if data.color is not None:
        c.color = data.color or None
    if data.style is not None:
        c.style = data.style
    db.commit()
    db.refresh(c)
    _recalc_progress(db, record.id)
    touch_record(db, record)
    return c


# alias for /cells/{id}
@router.put("/cells/{cell_id}", response_model=CellOut, include_in_schema=False)
def update_cell_alias(
    cell_id: str,
    data: CellUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    return update_cell(cell_id, data, db, user)


@router.post("/{record_id}/cells/bulk", response_model=list[CellOut])
def bulk_update(
    record_id: str,
    data: CellBulkUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    record = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not record:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    out = []
    for cell_id in data.ids:
        c = db.get(Cell, cell_id)
        if not c or c.record_id != record_id:
            continue
        if data.color is not None:
            c.color = data.color
        if data.style is not None:
            # merge style
            cur = c.style or {}
            cur.update(data.style)
            c.style = cur
        if data.reviewed is not None:
            c.reviewed = data.reviewed
        out.append(c)
    db.commit()
    for c in out:
        db.refresh(c)
    _recalc_progress(db, record_id)
    touch_record(db, record)
    return out


@router.get("/{record_id}/stats")
def stats(record_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    record = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not record:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    total = db.query(func.count(Cell.id)).filter(Cell.record_id == record_id).scalar() or 0
    rev = (
        db.query(func.count(Cell.id))
        .filter(Cell.record_id == record_id, Cell.reviewed.is_(True))
        .scalar()
        or 0
    )
    rows = _ensure_live(db, record)
    # count fully-reviewed rows (every row counts, with or without company)
    reviewed_count = 0
    for row in rows:
        t = (
            db.query(func.count(Cell.id))
            .filter(Cell.record_id == record_id, Cell.row_id == row.id)
            .scalar()
            or 0
        )
        r = (
            db.query(func.count(Cell.id))
            .filter(
                Cell.record_id == record_id,
                Cell.row_id == row.id,
                Cell.reviewed.is_(True),
            )
            .scalar()
            or 0
        )
        if t > 0 and t == r:
            reviewed_count += 1
    total_rows = len(rows)
    pending_count = total_rows - reviewed_count if total_rows else 0
    prog = int(rev * 100 / total) if total else 0
    return {
        "total": total_rows,
        "reviewed": reviewed_count,
        "pending": pending_count,
        "progress": prog,
        "total_cells": total,
        "reviewed_cells": rev,
    }


def _missing_fields(rows: list[RecordRow]) -> list[dict]:
    """Required fields for save/export: every row needs a registered company."""
    ordered = sorted(rows, key=lambda f: f.position)
    return find_rows_missing_company(
        [
            {"id": r.id, "company_id": r.company_id, "name_snapshot": r.name_snapshot}
            for r in ordered
        ]
    )


@router.get("/{record_id}/export")
def export_file(
    record_id: str,
    format: str = "pdf",
    lang: str = "es",
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    TXT = {
        "en": {
            "months": ["J", "F", "M", "A", "M", "J", "J", "A", "S", "O", "N", "D"],
            "number": "No.",
            "company": "Company",
            "assignee": "Assignee",
            "notes": "Notes",
            "sheet": "Review",
            "untitled": "Untitled review",
            "companies": "Companies",
            "file": "review",
        },
        "es": {
            "months": ["E", "F", "M", "A", "M", "J", "J", "A", "S", "O", "N", "D"],
            "number": "N.º",
            "company": "Empresa",
            "assignee": "Responsable",
            "notes": "Observaciones",
            "sheet": "Revisión",
            "untitled": "Revisión sin título",
            "companies": "Empresas",
            "file": "revision",
        },
    }["en" if lang == "en" else "es"]
    record = (
        db.query(Record)
        .filter(Record.id == record_id, Record.user_id == user.id, Record.deleted_at.is_(None))
        .first()
    )
    if not record:
        raise HTTPException(404, {"code": "RECORD_NOT_FOUND", "message": "Not found"})
    rows = _ensure_live(db, record)
    rows = sorted(rows, key=lambda f: f.position)
    missing = _missing_fields(rows)
    if missing:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "EXPORT_MISSING_FIELDS",
                "message": "Missing required fields to export",
                "faltantes": missing,
            },
        )
    cells = db.query(Cell).filter(Cell.record_id == record_id).all()
    cmap = {(c.row_id, c.year, c.month): c for c in cells}

    if format == "excel":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        wb = Workbook()
        ws = wb.active
        ws.title = TXT["sheet"]
        header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=9)
        thin = Side(style="thin", color="E2E8F0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        years = list(range(record.period_start, record.period_end + 1))
        months = TXT["months"]
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
        ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
        ws.cell(row=1, column=1, value=TXT["number"]).font = header_font
        ws.cell(row=1, column=1).fill = header_fill
        ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=1, column=2, value=TXT["company"]).font = header_font
        ws.cell(row=1, column=2).fill = header_fill
        ws.cell(row=1, column=2).alignment = Alignment(horizontal="center", vertical="center")
        col = 3
        for y in years:
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 11)
            c = ws.cell(row=1, column=col, value=str(y))
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center")
            col += 12
        col = 3
        for _ in years:
            for m in months:
                c = ws.cell(row=2, column=col, value=m)
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(horizontal="center")
                col += 1
        for hdr in (TXT["assignee"], TXT["notes"]):
            ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
            c = ws.cell(row=1, column=col, value=hdr)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            col += 1
        r = 3
        for idx, row in enumerate(rows, 1):
            ws.cell(row=r, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=2, value=row.name_snapshot)
            col = 3
            for y in years:
                for m_idx in range(1, 13):
                    c = cmap.get((row.id, y, m_idx))
                    v = "☑" if c and c.reviewed else "☐"
                    cell = ws.cell(row=r, column=col, value=v)
                    cell.alignment = Alignment(horizontal="center")
                    if c and c.reviewed:
                        colr = (c.color or "6366F1").lstrip("#")
                        cell.font = Font(color=colr)
                        # style bold
                        if c.style and c.style.get("bold"):
                            cell.font = Font(color=colr, bold=True)
                    col += 1
            ws.cell(row=r, column=col, value=row.assignee or "")
            ws.cell(row=r, column=col + 1, value=row.note or "")
            col += 2
            r += 1
        for row in ws.iter_rows(min_row=1, max_row=r - 1, max_col=col - 1):
            for cell in row:
                cell.border = border
        ws.freeze_panes = "C3"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return StreamingResponse(
            bio,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": (
                    f"attachment; filename={export_filename(record.title, 'xlsx')}"
                )
            },
        )
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        bio = io.BytesIO()
        doc = SimpleDocTemplate(
            bio,
            pagesize=landscape(A4),
            leftMargin=10 * mm,
            rightMargin=10 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm,
        )
        styles = getSampleStyleSheet()
        story = []
        doc_title = (
            f"{record.title or TXT['untitled']} — "
            f"{record.period_start}-{record.period_end} | "
            f"{len(rows)} {TXT['companies']}"
        )
        story.append(
            Paragraph(
                doc_title,
                styles["Title"],
            )
        )
        story.append(Spacer(1, 6))
        years = list(range(record.period_start, record.period_end + 1))
        months = TXT["months"]
        flat_header = [TXT["number"], TXT["company"]]
        for y in years:
            for m in months:
                flat_header.append(f"{y}-{m}")
        flat_header += [TXT["assignee"], TXT["notes"]]
        table_data = [flat_header]
        for idx, row in enumerate(rows, 1):
            entry = [str(idx), row.name_snapshot]
            for y in years:
                for m_idx in range(1, 13):
                    c = cmap.get((row.id, y, m_idx))
                    entry.append("☑" if c and c.reviewed else "☐")
            entry += [row.assignee or "", row.note or ""]
            table_data.append(entry)
        n_meses = len(flat_header) - 4
        col_widths = [12 * mm, 30 * mm] + [8 * mm] * n_meses + [20 * mm, 28 * mm]
        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        style = TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6366f1")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, 0), 6),
                ("FONTSIZE", (0, 1), (-1, -1), 5),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ]
        )
        t.setStyle(style)
        story.append(t)
        # Fase 5 §12: adjuntos de imagen como miniaturas al final del PDF.
        from pathlib import Path as _Path

        from reportlab.platypus import Image as _RLImage

        from app.core.config import settings as _settings
        from app.models.attachment import Attachment as _Attachment

        for att in db.query(_Attachment).filter(_Attachment.record_id == record.id).all():
            if not att.mime.startswith("image/"):
                continue
            for ext in (".jpg", ".png", ".webp"):
                p = _Path(_settings.ATTACH_DIR) / record.id / f"{att.hash}{ext}"
                if not p.exists():
                    continue
                try:
                    from PIL import Image as _PILImage

                    with _PILImage.open(p) as im:
                        im.verify()
                    with _PILImage.open(p) as im:
                        w, h = im.size
                    side = 30 * mm
                    scale = side / max(w, h)
                    story.append(Spacer(1, 6))
                    story.append(_RLImage(str(p), width=w * scale, height=h * scale))
                except Exception:
                    pass
                break
        doc.build(story)
        bio.seek(0)
        return StreamingResponse(
            bio,
            media_type="application/pdf",
            headers={
                "Content-Disposition": (
                    f"attachment; filename={export_filename(record.title, 'pdf')}"
                )
            },
        )
