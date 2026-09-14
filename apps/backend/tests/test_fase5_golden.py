"""Fase 5: cross-platform golden — server Excel == expected grid."""

import io
import json
from pathlib import Path

from openpyxl import load_workbook

from tests.conftest import auth_headers, client

# Fixtures golden cross-platform: antes vivian en packages/sync (paquete RN
# removed); now backend-local, its only consumer.
FIX = Path(__file__).resolve().parent / "fixtures"
GOLDEN_DOC = json.loads((FIX / "golden-doc.json").read_text())["doc"]
GOLDEN_EXCEL = json.loads((FIX / "golden-excel.json").read_text())


def test_server_excel_matches_golden_grid():
    h = auth_headers("gold@t.com")
    r = client.post("/api/sync/push", json={"changes": [GOLDEN_DOC]}, headers=h)
    assert r.json()["accepted"], r.text
    rid = client.get("/api/records", headers=h).json()[0]["id"]
    xls = client.get(f"/api/records/{rid}/export?format=excel", headers=h)
    assert xls.status_code == 200
    wb = load_workbook(io.BytesIO(xls.content), read_only=True)
    ws = wb.active
    assert ws is not None
    rows = [[c.value for c in row] for row in ws.iter_rows()]

    # openpyxl devuelve None en celdas vacias/fusionadas (""); equivalencia.
    def norm(grid):
        return [[("" if v is None else v) for v in r] for r in grid]

    assert norm(rows) == norm(GOLDEN_EXCEL["rows"])
